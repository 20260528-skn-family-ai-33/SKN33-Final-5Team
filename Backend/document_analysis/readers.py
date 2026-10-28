"""Read source objects without flattening Office documents to PDF.

OOXML is inspected even when a high-level library cannot expose an object. Only
package-local relationships are read; linked URLs are never fetched.
"""
import csv
import io
import posixpath
import re
import warnings as image_warnings
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from xml.etree import ElementTree as ET
from .contracts import AnalysisError, IMAGE_FORMATS, MAX_UNITS, MAX_BLOCKS, block, bounds

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      's': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}


def local(tag):
    return tag.rsplit('}', 1)[-1]


def text_of(node):
    chunks = []
    def visit(item):
        tag = local(item.tag)
        if tag == 't':
            chunks.append(item.text or '')
        elif tag == 'tab':
            chunks.append('\t')
        elif tag in {'br', 'cr'}:
            chunks.append('\n')
        else:
            for child in item:
                visit(child)
            if tag == 'p':
                chunks.append('\n')
    visit(node)
    return ''.join(chunks).rstrip('\n')


class Package:
    def __init__(self, source):
        try:
            self.zip = ZipFile(source)
            names = set()
            size = 0
            for entry in self.zip.infolist():
                if (entry.filename in names or '\\' in entry.filename or
                        entry.filename.startswith('/') or '..' in entry.filename.split('/') or entry.flag_bits & 1):
                    raise AnalysisError('invalid_package', 'Office 압축 파일의 항목이 올바르지 않습니다.')
                names.add(entry.filename)
                size += entry.file_size
                if entry.file_size > 16 * 1024 * 1024 or size > 64 * 1024 * 1024 or len(names) > 2000:
                    raise AnalysisError('package_limit', 'Office 압축 해제 크기 제한을 넘었습니다.')
            self.names = names
        except BadZipFile as exc:
            raise AnalysisError('invalid_package', '올바른 Office 압축 파일이 아닙니다.') from exc

    def close(self):
        self.zip.close()

    def xml(self, name):
        data = self.zip.read(name)
        if b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():
            raise AnalysisError('invalid_xml', '외부 엔터티가 있는 문서는 처리할 수 없습니다.')
        try:
            return ET.fromstring(data)
        except ET.ParseError as exc:
            raise AnalysisError('invalid_xml', '문서 XML을 읽을 수 없습니다.') from exc

    def relationships(self, part):
        rel = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')
        if rel not in self.names:
            return {}
        result = {}
        for item in self.xml(rel):
            target = item.get('Target', '')
            external = item.get('TargetMode') == 'External'
            if not external:
                target = posixpath.normpath(target.lstrip('/') if target.startswith('/') else
                                            posixpath.join(posixpath.dirname(part), target))
            if target.startswith('../'):
                raise AnalysisError('invalid_relationship', '문서 내부 객체 경로가 올바르지 않습니다.')
            result[item.get('Id')] = dict(target=target, external=external, type=item.get('Type', '').rsplit('/', 1)[-1])
        return result


class Inspector:
    def __init__(self, source, work):
        self.source = Path(source)
        self.work = Path(work)
        self.units = []
        self.blocks = []
        self.warnings = []
        self.images = []
        self.signals = []
        self.package = None

    def add(self, kind, location, **kwargs):
        if len(self.blocks) >= MAX_BLOCKS:
            raise AnalysisError('object_limit', '문서 객체 수 제한을 넘었습니다.')
        item = block('native-' + str(len(self.blocks) + 1), kind, location, **kwargs)
        self.blocks.append(item)
        return item

    def unit(self, identifier, location):
        if len(self.units) >= MAX_UNITS:
            raise AnalysisError('unit_limit', '문서는 최대 500개 처리 단위로 나누세요.')
        unit = dict(id=identifier, native_location=location)
        self.units.append(unit)
        return identifier

    def image(self, data, suffix, location, *, parent=None, unit=None, bbox=None):
        folder = self.work / 'images'
        folder.mkdir(exist_ok=True, parents=True)
        path = folder / f'image-{len(self.images) + 1}{suffix.lower()}'
        path.write_bytes(data)
        item = self.add('image', location, parent=parent, unit=unit, bbox=bbox,
                        artifact=str(path.relative_to(self.work)))
        self.images.append(dict(block_id=item['id'], path=path))
        return item

    def related_objects(self, node, part, location, unit, parent=None):
        relations = self.package.relationships(part)
        for entry in node.iter():
            if local(entry.tag) not in {'blip', 'chart', 'imagedata', 'OLEObject'}:
                continue
            rel_id = entry.get('{' + NS['r'] + '}embed') or entry.get('{' + NS['r'] + '}id') or entry.get('{' + NS['r'] + '}link')
            rel = relations.get(rel_id)
            if not rel:
                self.warnings.append('unresolved_object_relationship')
                continue
            loc = {**location, 'package_part': part, 'relationship_id': rel_id}
            if rel['external']:
                self.add('external_reference', loc, unit=unit, parent=parent, data={'target': rel['target']}, status='review_required')
                self.warnings.append('external_reference_not_fetched')
            elif rel['type'] == 'image' and rel['target'] in self.package.names:
                self.image(self.package.zip.read(rel['target']), Path(rel['target']).suffix, loc, parent=parent, unit=unit)
            elif rel['type'] == 'chart' and rel['target'] in self.package.names:
                chart = self.package.xml(rel['target'])
                values = [x.text for x in chart.iter() if local(x.tag) in {'v', 'f'} and x.text]
                self.add('chart', loc, unit=unit, parent=parent, data={'source_part': rel['target'], 'cached_values_and_references': values},
                         status='review_required')
                self.warnings.append('chart_semantics_not_verified')
            else:
                self.add('unknown', loc, unit=unit, parent=parent, status='review_required')
                self.warnings.append('unsupported_embedded_object')

    def office(self, fmt):
        self.package = Package(self.source)
        try:
            if fmt == 'docx':
                self.word()
            elif fmt == 'pptx':
                self.slides()
            else:
                self.spreadsheet()
        finally:
            self.package.close()
        # Read-only high-level inspection supplements package relationships; it
        # never resaves the source or replaces formula/cache XML evidence.
        try:
            if fmt == 'docx':
                from docx import Document
                doc = Document(self.source)
                self.signals.append(dict(reader='python-docx', paragraphs=len(doc.paragraphs), tables=len(doc.tables), inline_shapes=len(doc.inline_shapes)))
            elif fmt == 'pptx':
                from pptx import Presentation
                doc = Presentation(self.source)
                self.signals.append(dict(reader='python-pptx', slides=len(doc.slides)))
        except ImportError:
            self.warnings.append('office_high_level_reader_unavailable_xml_only')

    def word(self):
        part = 'word/document.xml'
        root = self.package.xml(part)
        unit = self.unit('document', {'document': True})
        body = root.find('w:body', NS)
        if body is None:
            raise AnalysisError('invalid_docx', 'Word 본문을 찾을 수 없습니다.')
        for index, item in enumerate(body, 1):
            kind = local(item.tag)
            location = {'body_index': index, 'object_type': kind}
            if kind == 'p':
                parent = self.add('text', location, text=text_of(item), unit=unit)
                self.related_objects(item, part, location, unit, parent['id'])
                if item.findall('.//w:txbxContent', NS):
                    parent['status'] = 'review_required'
                    self.warnings.append('floating_textbox_order_requires_review')
            elif kind == 'tbl':
                cells = []
                for row_index, row in enumerate(item.findall('w:tr', NS)):
                    col = 0
                    for cell in row.findall('w:tc', NS):
                        span = cell.find('w:tcPr/w:gridSpan', NS)
                        merge = cell.find('w:tcPr/w:vMerge', NS)
                        col_span = int(span.get('{' + NS['w'] + '}val', '1')) if span is not None else 1
                        cells.append(dict(row=row_index, column=col, col_span=col_span,
                                          vertical_merge=merge.get('{' + NS['w'] + '}val', 'continue') if merge is not None else None,
                                          text=text_of(cell)))
                        col += col_span
                parent = self.add('table', location, unit=unit, text='\n'.join(c['text'] for c in cells), data={'cells': cells})
                self.related_objects(item, part, location, unit, parent['id'])
            elif kind != 'sectPr':
                self.add('unknown', location, unit=unit, status='review_required')
                self.warnings.append('unprocessed_word_object')
        for name in self.package.names:
            if re.match(r'word/(header|footer|footnotes|endnotes|comments).*\.xml$', name):
                self.warnings.append('word_auxiliary_part_not_in_body:' + name)

    def slides(self):
        root = self.package.xml('ppt/presentation.xml')
        relations = self.package.relationships('ppt/presentation.xml')
        for index, entry in enumerate(root.findall('p:sldIdLst/p:sldId', NS), 1):
            rel = relations.get(entry.get('{' + NS['r'] + '}id'))
            if not rel or rel['external']:
                raise AnalysisError('invalid_pptx', '슬라이드 연결을 확인하세요.')
            part = rel['target']
            unit = self.unit('slide-' + str(index), {'slide': index})
            root_slide = self.package.xml(part)
            tree = root_slide.find('p:cSld/p:spTree', NS)
            if tree is None:
                self.warnings.append('empty_slide_tree')
                continue
            def shapes(node, parent=None):
                for item in node:
                    kind = local(item.tag)
                    if kind in {'nvGrpSpPr', 'grpSpPr'}:
                        continue
                    identity = next((x for x in item.iter() if local(x.tag) == 'cNvPr'), None)
                    location = {'slide': index, 'shape_id': identity.get('id') if identity is not None else None}
                    table = item.find('.//a:tbl', NS)
                    label = 'group' if kind == 'grpSp' else 'table' if table is not None else 'text'
                    data = {}
                    if table is not None:
                        data['cells'] = [dict(row=r, column=c, text=text_of(cell), row_span=cell.get('rowSpan'), col_span=cell.get('gridSpan'))
                                         for r, row in enumerate(table.findall('a:tr', NS)) for c, cell in enumerate(row.findall('a:tc', NS))]
                    transform = next((x for x in item.iter() if local(x.tag) == 'xfrm'), None)
                    if transform is not None:
                        data['native_transform_xml'] = ET.tostring(transform, encoding='unicode')
                    shape = self.add(label, location, unit=unit, parent=parent, text='' if kind == 'grpSp' else text_of(item), data=data)
                    if kind == 'grpSp':
                        shapes(item, shape['id'])
                    else:
                        self.related_objects(item, part, location, unit, shape['id'])
                        if kind not in {'sp', 'pic', 'graphicFrame', 'cxnSp'}:
                            shape['status'] = 'review_required'
                            self.warnings.append('unprocessed_slide_object')
            shapes(tree)
            if any(r['type'] == 'notesSlide' for r in self.package.relationships(part).values()):
                self.warnings.append('speaker_notes_not_included:slide-' + str(index))

    def spreadsheet(self):
        package = self.package
        root = package.xml('xl/workbook.xml')
        relations = package.relationships('xl/workbook.xml')
        strings = []
        if 'xl/sharedStrings.xml' in package.names:
            strings = [text_of(x) for x in package.xml('xl/sharedStrings.xml')]
        formats = {}
        styles = []
        if 'xl/styles.xml' in package.names:
            style = package.xml('xl/styles.xml')
            formats = {int(x.get('numFmtId')): x.get('formatCode') for x in style.findall('s:numFmts/s:numFmt', NS)}
            styles = [int(x.get('numFmtId', '0')) for x in style.findall('s:cellXfs/s:xf', NS)]
        for sheet in root.findall('s:sheets/s:sheet', NS):
            name = sheet.get('name')
            rel = relations.get(sheet.get('{' + NS['r'] + '}id'))
            if not rel or rel['external']:
                raise AnalysisError('invalid_xlsx', '시트 연결을 확인하세요.')
            unit = self.unit('sheet-' + str(len(self.units) + 1), {'sheet': name})
            sheet_xml = package.xml(rel['target'])
            merges = [x.get('ref') for x in sheet_xml.findall('s:mergeCells/s:mergeCell', NS)]
            parent = self.add('sheet', {'sheet': name}, unit=unit, data={'state': sheet.get('state', 'visible'), 'merged_ranges': merges})
            for cell in sheet_xml.findall('s:sheetData/s:row/s:c', NS):
                address, value_type = cell.get('r'), cell.get('t', 'n')
                value, formula = cell.find('s:v', NS), cell.find('s:f', NS)
                raw = value.text if value is not None else None
                native_value = strings[int(raw)] if value_type == 's' and raw is not None else text_of(cell) if value_type == 'inlineStr' else raw
                style_index = int(cell.get('s', '0'))
                format_id = styles[style_index] if style_index < len(styles) else 0
                data = dict(raw_value=raw, value=native_value, value_type=value_type, number_format_id=format_id,
                            number_format=formats.get(format_id), formula=formula.text if formula is not None else None,
                            formula_attributes=dict(formula.attrib) if formula is not None else {},
                            cached_value=raw if formula is not None else None)
                if formula is not None and raw is None:
                    self.warnings.append('formula_cache_missing:' + name + '!' + str(address))
                content = str(native_value) if native_value is not None else '=' + (formula.text or '') if formula is not None else ''
                self.add('cell', {'sheet': name, 'cell': address}, unit=unit, parent=parent['id'], text=str(address) + ': ' + content, data=data)
            for drawing in sheet_xml.findall('s:drawing', NS):
                drawing_rel = package.relationships(rel['target']).get(drawing.get('{' + NS['r'] + '}id'))
                if not drawing_rel or drawing_rel['external']:
                    self.warnings.append('unresolved_spreadsheet_drawing')
                    continue
                for anchor in package.xml(drawing_rel['target']):
                    location = {'sheet': name, 'anchor_xml': ET.tostring(anchor, encoding='unicode')[:4000]}
                    self.related_objects(anchor, drawing_rel['target'], location, unit, parent['id'])
            if sheet_xml.find('s:legacyDrawing', NS) is not None:
                self.warnings.append('legacy_spreadsheet_drawing_unprocessed:' + name)
        # XML preserves both <f> and <v>; openpyxl supplements the typed/raw view.
        try:
            import openpyxl
            book = openpyxl.load_workbook(self.source, read_only=True, data_only=False, keep_links=False)
            try:
                for item in self.blocks:
                    if 'cell' in item['content_types']:
                        cell = book[item['native_location']['sheet']][item['native_location']['cell']]
                        item['data']['number_format'] = cell.number_format
                        item['data']['openpyxl_data_type'] = cell.data_type
                        item['engine'] = 'openpyxl+xml'
            finally:
                book.close()
        except ImportError:
            self.warnings.append('openpyxl_unavailable_xml_only')

    def pdf(self):
        try:
            try:
                import pymupdf as fitz
            except ImportError:
                import fitz
        except ImportError as exc:
            raise AnalysisError('dependency_missing', 'PDF 구조 조사에 PyMuPDF가 필요합니다.') from exc
        with fitz.open(self.source) as doc:
            if doc.needs_pass:
                raise AnalysisError('encrypted_pdf', '암호화된 PDF는 암호를 해제한 뒤 연결하세요.')
            if len(doc) > MAX_UNITS:
                raise AnalysisError('unit_limit', 'PDF는 최대 500페이지로 나누세요.')
            for index, page in enumerate(doc, 1):
                unit = self.unit('page-' + str(index), {'page': index})
                text = page.get_text()
                images = page.get_image_info()
                drawings = page.get_drawings()
                area = max(page.rect.width * page.rect.height, 1)
                # Union area is not assumed; overlap is capped and this remains a heuristic.
                image_fraction = min(1, sum(max(0, x['bbox'][2]-x['bbox'][0]) * max(0, x['bbox'][3]-x['bbox'][1]) for x in images) / area)
                letters = sum(c.isalnum() for c in text)
                broken = sum(c == '\ufffd' or ord(c) < 32 and not c.isspace() for c in text) / max(len(text), 1)
                self.signals.append(dict(unit_id=unit, alphanumeric_count=letters, broken_ratio=broken,
                                         image_area_fraction_upper_bound=image_fraction, image_count=len(images),
                                         vector_drawing_count=len(drawings),
                                         width=page.rect.width, height=page.rect.height,
                                         scan_candidate=letters < 30 and image_fraction >= .6,
                                         text_layer_candidate=letters >= 30 and broken < .05))
                for item in page.get_text('blocks'):
                    if item[6] == 0:
                        b = self.add('text', {'page': index}, unit=unit, text=item[4], engine='pymupdf',
                                     bbox=bounds(list(item[:4]), 'pdf_top_left_points', width=page.rect.width, height=page.rect.height))
                        b['selected'] = False  # Evidence, not a second copy of parser output.
                for drawing in drawings:
                    rect = drawing.get('rect')
                    if rect is not None:
                        item = self.add('vector_graphic', {'page': index}, unit=unit, engine='pymupdf',
                                        bbox=bounds(list(rect), 'pdf_top_left_points', width=page.rect.width, height=page.rect.height),
                                        data={'drawing_type': drawing.get('type'), 'sequence_number': drawing.get('seqno')})
                        item['selected'] = False

    def image_file(self):
        try:
            from PIL import Image, ImageOps
            if self.source.suffix.lower() == '.heic':
                import pillow_heif
                pillow_heif.register_heif_opener()
        except ImportError as exc:
            raise AnalysisError('dependency_missing', '이미지 처리에 Pillow / HEIC에는 pillow-heif가 필요합니다.') from exc
        with image_warnings.catch_warnings():
            image_warnings.simplefilter('error', Image.DecompressionBombWarning)
            with Image.open(self.source) as image:
                count = getattr(image, 'n_frames', 1)
                if count > MAX_UNITS:
                    raise AnalysisError('unit_limit', '이미지는 최대 500프레임으로 나누세요.')
                for frame in range(count):
                    image.seek(frame)
                    orientation = image.getexif().get(274, 1)
                    original_size = image.size
                    upright = ImageOps.exif_transpose(image.copy()).convert('RGB')
                    unit = self.unit('frame-' + str(frame + 1), {'frame': frame + 1})
                    stream = io.BytesIO()
                    upright.save(stream, format='PNG')
                    item = self.image(stream.getvalue(), '.png', {'frame': frame + 1}, unit=unit,
                                      bbox=bounds([0, 0, *upright.size], 'upright_image_pixels', width=upright.width, height=upright.height))
                    item['data']['transform'] = dict(exif_orientation=orientation, original_size=list(original_size),
                                                     upright_size=list(upright.size), operation='exif_transpose')

    def plain(self, fmt):
        raw = self.source.read_bytes()
        try:
            text = raw.decode('utf-8-sig')
            encoding = 'utf-8-sig'
        except UnicodeDecodeError:
            try:
                text = raw.decode('cp949')
                encoding = 'cp949'
            except UnicodeDecodeError as exc:
                raise AnalysisError('unsupported_encoding', '텍스트 인코딩을 UTF-8 또는 CP949로 저장하세요.') from exc
        unit = self.unit('document', {'document': True})
        if fmt in {'csv', 'tsv'}:
            try:
                reader = csv.reader(io.StringIO(text, newline=''), delimiter='\t' if fmt == 'tsv' else ',', strict=True)
                for index, row in enumerate(reader, 1):
                    self.add('table_row', {'logical_row': index}, unit=unit, text='\t'.join(row),
                             data={'fields': row, 'encoding': encoding, 'physical_line_end': reader.line_num})
            except csv.Error as exc:
                raise AnalysisError('invalid_delimited_text', '구분자·인용부호를 확인하세요.') from exc
        else:
            self.add('text', {'document': True}, unit=unit, text=text, data={'encoding': encoding})
            if fmt == 'md':
                refs = re.findall(r'!\[([^\]]*)\]\(([^)]+)\)', text)
                for caption, target in refs:
                    self.add('external_reference', {'document': True}, unit=unit, data={'caption': caption, 'target': target}, status='review_required')
                if refs:
                    self.warnings.append('markdown_image_references_not_fetched')

    def run(self, fmt):
        if fmt in {'docx', 'pptx', 'xlsx'}:
            self.office(fmt)
        elif fmt == 'pdf':
            self.pdf()
        elif fmt in IMAGE_FORMATS:
            self.image_file()
        else:
            self.plain(fmt)
        return dict(units=self.units, blocks=self.blocks, images=self.images,
                    warnings=list(dict.fromkeys(self.warnings)), signals=self.signals)
