"""Real source fixtures; model doubles test routing, not recognition accuracy."""
import io
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from PIL import Image
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from .contracts import AnalysisError, validate, sha256
from .pipeline import run
from .readers import Package


class Engines:
    def __init__(self, fail=(), *, table=False):
        self.fail = set(fail)
        self.table = table
        self.attempts = []

    def run(self, engine, *, source=None, sources=None, scope=None, phase='primary'):
        attempt = dict(id='attempt-' + str(len(self.attempts)), processor=engine, phase=phase,
                       scope=scope, elapsed_seconds=.01, status='failed' if engine in self.fail else 'succeeded')
        self.attempts.append(attempt)
        if engine in self.fail:
            raise AnalysisError('dependency_missing', 'model not installed')
        base = dict(engine=engine, version='test-double', model={'name': 'test-double'}, partial=False)
        if engine == 'layout':
            return {**base, 'outputs': [dict(source=str(p), result={'boxes': [dict(label='table' if self.table else 'text', score=.98, coordinate=[2, 3, 95, 80])]}) for p in sources]}
        if engine == 'paddleocr':
            return {**base, 'outputs': [dict(source=str(p), result={'rec_texts': ['정산금엑 000017', '휴대전/화'], 'rec_scores': [.95, .96]}) for p in sources]}
        return {**base, 'page_count': 1, 'blocks': [dict(kind='table' if self.table else 'text', text='parsed 000017',
                location={'page': 1}, bbox=[0, 0, .8, .8] if engine == 'mineru' else None,
                coordinate_system='mineru_normalized_top_left', data={'cells': [{'row': 0, 'column': 0, 'text': '000017'}]} if self.table else {})]}


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.work = self.root / 'work'

    def tearDown(self):
        self.temp.cleanup()

    def picture(self, name='receipt.png'):
        path = self.root / name
        Image.new('RGB', (100, 90), 'white').save(path)
        return path

    def spreadsheet(self, image=True):
        source = self.root / 'mixed.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.title = '정산'
        sheet.append(['거래번호', '금액'])
        sheet.append(['000017', 45000])
        sheet['A3'] = 17
        sheet['A3'].number_format = '000000'
        sheet['B3'] = '=SUM(B2:B2)'
        sheet.merge_cells('A4:B4')
        sheet['A4'] = '병합 값'
        if image:
            sheet.add_image(ExcelImage(io.BytesIO(self.picture().read_bytes())), 'D2')
        hidden = book.create_sheet('숨김')
        hidden.sheet_state = 'hidden'
        hidden['A1'] = '다른 시트'
        book.save(source)
        return source

    def test_excel_native_values_formulas_hidden_sheets_and_ocr_provenance(self):
        source = self.spreadsheet()
        before = sha256(source)
        engine = Engines()
        payload = run(source, self.work, engines=engine)
        data = payload['document_analysis']
        cells = {b['native_location'].get('cell'): b for b in data['blocks'] if b['content_types'] == ['cell'] and b['native_location']['sheet'] == '정산'}
        self.assertEqual(cells['A2']['data']['value'], '000017')
        self.assertEqual(cells['A3']['data']['value'], '17')
        self.assertEqual(cells['A3']['data']['number_format'], '000000')
        self.assertEqual(cells['B3']['data']['formula'], 'SUM(B2:B2)')
        self.assertIsNone(cells['B3']['data']['cached_value'])
        self.assertEqual(len(data['units']), 2)
        sheet = next(b for b in data['blocks'] if b['content_types'] == ['sheet'] and b['native_location']['sheet'] == '정산')
        self.assertEqual(sheet['data']['merged_ranges'], ['A4:B4'])
        ocr = next(b for b in data['blocks'] if b['engine'] == 'paddleocr')
        parent = next(b for b in data['blocks'] if b['id'] == ocr['parent_object_id'])
        self.assertEqual(parent['native_location']['sheet'], '정산')
        self.assertEqual(parent['bbox']['coordinate_system'], 'upright_image_pixels')
        self.assertEqual(ocr['data']['crop_origin'], [2, 3])
        self.assertIn('정산금엑 000017', payload['text'])
        self.assertEqual([a['processor'] for a in engine.attempts], ['layout', 'paddleocr'])
        self.assertEqual(sha256(source), before)
        validate(data)

    def test_missing_layout_is_partial_and_keeps_image_and_cells(self):
        payload = run(self.spreadsheet(), self.work, engines=Engines(fail={'layout'}))
        data = payload['document_analysis']
        self.assertEqual(payload['parser_status'], 'partial')
        self.assertIn('image_classification:dependency_missing', data['review_reasons'])
        self.assertEqual(data['coverage']['unprocessed_image_objects'], 1)
        self.assertTrue(any(b['artifact'] for b in data['blocks'] if 'image' in b['content_types']))
        self.assertIn('000017', payload['text'])

    def test_image_table_goes_to_docling_and_crop_coordinates_stay_local(self):
        payload = run(self.picture(), self.work, engines=Engines(table=True))
        data = payload['document_analysis']
        self.assertEqual([r['processor'] for r in data['routes']], ['layout', 'docling'])
        self.assertEqual(len(payload['tables']), 1)
        table = next(b for b in data['blocks'] if b['engine'] == 'docling')
        self.assertIsNotNone(table['parent_object_id'])
        self.assertEqual(table['native_location'], {'frame': 1})

    def test_known_dpi_survives_normalization_rotation_and_table_crop(self):
        for orientation in (1, 6):
            with self.subTest(orientation=orientation):
                source = self.root / f'dpi-{orientation}.png'
                exif = Image.Exif()
                exif[274] = orientation
                Image.new('RGB', (100, 90), 'white').save(source, dpi=(144, 288), exif=exif)
                before = sha256(source)
                work = self.root / f'dpi-work-{orientation}'
                payload = run(source, work, engines=Engines(table=True))
                expected = (288, 144) if orientation == 6 else (144, 288)
                artifacts = list((work / 'images').glob('*.png')) + list(work.glob('upright-*.png')) + list(work.glob('crop-*.png'))
                self.assertEqual(len(artifacts), 3)
                for artifact in artifacts:
                    with Image.open(artifact) as image:
                        for actual, gold in zip(image.info['dpi'], expected):
                            self.assertAlmostEqual(actual, gold, delta=.03)
                region = next(b for b in payload['document_analysis']['blocks'] if b['engine'] == 'layout')
                self.assertEqual(region['data']['crop_origin'], [2, 3])
                self.assertEqual(sha256(source), before)

    def test_missing_crop_table_is_compared_with_image_within_supplement_budget(self):
        class CropMissEngines(Engines):
            def run(self, engine, **kwargs):
                parsed = super().run(engine, **kwargs)
                if engine == 'docling' and kwargs.get('phase') == 'supplement':
                    parsed['blocks'][0].update(kind='text', data={})
                return parsed
        for budget in (1, 2):
            with self.subTest(budget=budget):
                engine = CropMissEngines(table=True)
                payload = run(self.picture(), self.root / f'budget-{budget}', engines=engine, max_supplements=budget)
                analysis = payload['document_analysis']
                self.assertEqual(payload['parser_status'], 'partial')
                self.assertEqual(len(engine.attempts), 2 if budget == 1 else 3)
                candidates = [b for b in analysis['blocks'] if b.get('phase') == 'comparison']
                if budget == 1:
                    self.assertEqual(candidates, [])
                    self.assertIn('supplement_limit_reached', analysis['review_reasons'])
                else:
                    self.assertEqual(len(candidates), 1)
                    candidate = candidates[0]
                    self.assertFalse(candidate['selected'])
                    self.assertEqual(candidate['parent_object_id'], analysis['blocks'][0]['id'])
                    self.assertEqual(candidate['data']['context_for_regions'], [analysis['blocks'][1]['id']])
                    self.assertEqual(candidate['data']['input_coordinate_system'], 'containing_upright_image')
                    self.assertNotIn('crop_origin', candidate['data'])
                    self.assertEqual(engine.attempts[-1]['phase'], 'comparison')

    def test_unknown_or_invalid_dpi_is_not_invented(self):
        from .image_metadata import png_resolution
        image = Image.new('RGB', (100, 90), 'white')
        for value in (None, (float('nan'), 144), (0, 144), (1e300, 144), 'bad', (144,)):
            image.info['dpi'] = value
            self.assertEqual(png_resolution(image), {})
        run(self.picture(), self.work, engines=Engines(table=True))
        for path in self.work.glob('*.png'):
            with Image.open(path) as normalized:
                self.assertNotIn('dpi', normalized.info)

    def test_pdf_figure_raster_dpi_matches_render_scale(self):
        import pymupdf
        source = self.root / 'figure.pdf'
        with pymupdf.open() as doc:
            doc.new_page(width=100, height=100)
            doc.save(source)
        class FigureEngines(Engines):
            def run(self, engine, **kwargs):
                parsed = super().run(engine, **kwargs)
                if engine == 'mineru':
                    parsed['blocks'][0].update(kind='image', text='', data={})
                return parsed
        payload = run(source, self.work, engines=FigureEngines(table=True))
        figure = next(b for b in payload['document_analysis']['blocks'] if b['engine'] == 'mineru')
        transform = figure['data']['raster_transform']
        self.assertEqual(transform['raster_dpi'], [108, 108])
        self.assertEqual(transform['pixels_per_point'], 1.5)
        for pattern in ('figure-*.png', 'upright-*.png', 'crop-*.png'):
            paths = list(self.work.glob(pattern))
            self.assertEqual(len(paths), 1)
            with Image.open(paths[0]) as image:
                self.assertAlmostEqual(image.info['dpi'][0], 108, delta=.03)

    def test_multiframe_tiff_and_supplement_limit_are_not_silently_complete(self):
        source = self.root / 'frames.tiff'
        images = [Image.new('RGB', (100, 90), color) for color in ['white', 'gray', 'blue']]
        images[0].save(source, save_all=True, append_images=images[1:])
        payload = run(source, self.work, engines=Engines(), max_supplements=1)
        data = payload['document_analysis']
        self.assertEqual(data['coverage']['total_units'], 3)
        self.assertEqual(data['coverage']['source_image_objects'], 3)
        self.assertEqual(data['coverage']['unprocessed_image_objects'], 2)
        self.assertIn('supplement_limit_reached', data['review_reasons'])
        self.assertEqual(payload['parser_status'], 'partial')

    def test_word_order_table_breaks_and_inserted_picture_are_preserved(self):
        source = self.root / 'mixed.docx'
        doc = Document()
        doc.add_paragraph('처음 000017')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = '휴대전\n화'
        table.cell(0, 1).text = '[V] 동의함'
        doc.add_picture(str(self.picture()))
        doc.add_paragraph('마지막')
        doc.save(source)
        payload = run(source, self.work, engines=Engines())
        self.assertIn('휴대전\n화', payload['text'])
        self.assertLess(payload['text'].index('처음'), payload['text'].index('마지막'))
        self.assertTrue(any(b['content_types'] == ['image'] and b['parent_object_id'] for b in payload['document_analysis']['blocks']))
        self.assertTrue(all('page' not in b['native_location'] for b in payload['document_analysis']['blocks'] if b['engine'] == 'native'))

    def test_slide_locations_and_group_transforms_are_preserved(self):
        source = self.root / 'slides.pptx'
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        shape = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        shape.text = '그룹 내용'
        slide.shapes.add_picture(str(self.picture()), Inches(1), Inches(3))
        presentation.save(source)
        payload = run(source, self.work, engines=Engines())
        groups = [b for b in payload['document_analysis']['blocks'] if b['content_types'] == ['group']]
        child = next(b for b in payload['document_analysis']['blocks'] if b['text'] == '그룹 내용')
        self.assertEqual(child['parent_object_id'], groups[0]['id'])
        self.assertEqual(child['native_location']['slide'], 1)
        self.assertIn('native_transform_xml', child['data'])

    def test_csv_quoted_newline_and_leading_zero_use_no_models(self):
        source = self.root / 'data.csv'
        source.write_text('코드,설명\n000017,"줄1\n줄2"\n', encoding='utf-8')
        engine = Engines()
        payload = run(source, self.work, engines=engine)
        blocks = payload['document_analysis']['blocks']
        self.assertEqual(blocks[1]['data']['fields'], ['000017', '줄1\n줄2'])
        self.assertEqual(blocks[1]['native_location'], {'logical_row': 2})
        self.assertEqual(engine.attempts, [])

    def test_markdown_references_are_recorded_and_never_fetched(self):
        source = self.root / 'notes.md'
        source.write_text('# 원문\n![그림](https://example.invalid/private.png)', encoding='utf-8')
        payload = run(source, self.work, engines=Engines())
        self.assertEqual(payload['parser_status'], 'partial')
        self.assertIn('markdown_image_references_not_fetched', payload['document_analysis']['review_reasons'])

    def test_package_rejects_external_entities_and_path_traversal(self):
        source = self.root / 'bad.docx'
        with ZipFile(source, 'w') as archive:
            archive.writestr('../escape.xml', 'bad')
        with self.assertRaises(AnalysisError):
            Package(source)
        with ZipFile(source, 'w') as archive:
            archive.writestr('word/document.xml', '<!DOCTYPE x [<!ENTITY bad SYSTEM "file:///etc/passwd">]><x>&bad;</x>')
        package = Package(source)
        try:
            with self.assertRaises(AnalysisError):
                package.xml('word/document.xml')
        finally:
            package.close()

    def test_pdf_fallback_and_hwp_converted_page_are_explicit(self):
        import pymupdf
        pdf = self.root / 'source.pdf'
        with pymupdf.open() as doc:
            page = doc.new_page()
            page.insert_text((30, 50), 'original 000017')
            doc.save(pdf)
        payload = run(pdf, self.work, engines=Engines(fail={'mineru'}))
        self.assertEqual(payload['document_analysis']['fallback_processor'], 'docling')
        self.assertEqual(payload['parser_status'], 'partial')
        fake_hwp = self.root / 'source.hwp'
        fake_hwp.write_bytes(b'fake hwp only tests conversion contract')
        def converter(source, fmt, work):
            return pdf, 'pdf', dict(artifact='converted/source.pdf', source_sha256=sha256(source), converted_sha256=sha256(pdf))
        payload = run(fake_hwp, self.root/'hwp-work', engines=Engines(), converter=converter)
        for b in payload['document_analysis']['blocks']:
            self.assertNotIn('page', b['native_location'])
            self.assertEqual(b['native_location']['converted_page'], 1)

    def test_contract_rejects_broken_parent_and_cycle(self):
        payload = run(self.picture(), self.work, engines=Engines())
        data = payload['document_analysis']
        data['blocks'][0]['parent_object_id'] = data['blocks'][-1]['id']
        with self.assertRaises(AnalysisError):
            validate(data)

    def test_table_cells_preserve_spans_and_leading_zero(self):
        from .tables import from_html
        table = from_html('<table><tr><th rowspan="2">번호</th><td>000017</td></tr><tr><td colspan="2">45,000</td></tr></table>')
        self.assertEqual(table['cells'][0]['row_span'], 2)
        self.assertEqual(table['cells'][1]['text'], '000017')
        self.assertEqual(table['cells'][2]['column'], 1)
        self.assertEqual(table['cells'][2]['col_span'], 2)


if __name__ == '__main__':
    unittest.main()
