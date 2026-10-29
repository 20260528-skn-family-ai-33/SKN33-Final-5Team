"""Freeze public and independently generated gold data before any inference."""
import argparse
import collections
import csv
import hashlib
import json
from pathlib import Path
import random
import shutil
import urllib.parse
import urllib.request

ROOT = Path(__file__).resolve().parents[1]
REVISION = 'f5f559bddf50e36f7f9899d842d0006f13ce8afc'
SEED = 20261028


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def pdf_index():
    url = f'https://huggingface.co/api/datasets/opendatalab/OmniDocBench/tree/{REVISION}/pdfs?limit=1000'
    entries = []
    while url:
        with urllib.request.urlopen(url, timeout=60) as response:
            entries.extend(json.load(response))
            links = response.headers.get('Link', '')
        url = next((part.split('>')[0].strip()[1:] for part in links.split(',') if 'rel="next"' in part), None)
    return {Path(row['path']).name: row['path'] for row in entries if row['type'] == 'file'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=ROOT / 'data/document-benchmark/v1')
    args = parser.parse_args()
    output = args.output.resolve()
    manifest_path = output / 'manifest.json'
    if manifest_path.exists():
        raise RuntimeError('고정된 manifest가 이미 있습니다. 다른 출력 디렉터리를 지정하세요.')
    public = ROOT / 'data/bmt-public'
    gold_all = json.loads((public / 'OmniDocBench-v1_0.json').read_text())
    available = pdf_index()
    groups = collections.defaultdict(list)
    for row in gold_all:
        image_name = Path(row['page_info']['image_path']).name
        desired = image_name.rsplit('.', 1)[0] + '.pdf'
        matches = [p for name, p in available.items() if name == desired or name == 'docstructbench_' + desired]
        if len(matches) == 1:
            groups[row['page_info']['page_attribute']['data_source']].append((row, matches[0]))
    assert len(groups) == 9 and all(len(rows) >= 2 for rows in groups.values()), {k: len(v) for k, v in groups.items()}
    rng = random.Random(SEED)
    selected = []
    for category in sorted(groups):
        candidates = sorted(groups[category], key=lambda pair: pair[1])
        rng.shuffle(candidates)
        used = set()
        for row, path in candidates:
            original = path.rsplit('.pdf_', 1)[0]
            if original in used:
                continue
            selected.append((row, path))
            used.add(original)
            if len(used) == 2:
                break
    output.mkdir(parents=True, exist_ok=True)
    write_json(output / 'gold/omni-selected.json', [r for r, _ in selected])
    cases = []
    def add(identifier, source, suite, gold=None, **kwargs):
        case = dict(id=identifier, input=str(source.relative_to(output)), name=source.name, sha256=digest(source),
                    bytes=source.stat().st_size, format=source.suffix[1:].lower(), suite=suite, **kwargs)
        if gold:
            case.update(gold=str(gold.relative_to(output)), gold_sha256=digest(gold))
        cases.append(case)
        return case
    for index, (row, source_path) in enumerate(selected, 1):
        source = output / 'inputs/omni' / f'omni-{index:02d}.pdf'
        source.parent.mkdir(parents=True, exist_ok=True)
        url = f'https://huggingface.co/datasets/opendatalab/OmniDocBench/resolve/{REVISION}/' + urllib.parse.quote(source_path, safe='/')
        with urllib.request.urlopen(url, timeout=120) as response, source.open('wb') as target:
            shutil.copyfileobj(response, target)
        add(f'OMNI-{index:02d}', source, 'omnidocbench', output / 'gold/omni-selected.json',
            source_url=url, original_pdf_path=source_path, image_id=Path(row['page_info']['image_path']).name,
            category=row['page_info']['page_attribute']['data_source'])
    pollock = public / 'Pollock/survey_sample'
    original_files = sorted((pollock / 'csv').iterdir(), key=lambda p: p.name)
    assert len(original_files) == 100
    for index, source in enumerate(original_files, 1):
        gold = pollock / 'clean' / source.name
        assert gold.is_file()
        dest = output / 'inputs/pollock' / source.name
        target_gold = output / 'gold/pollock' / source.name
        dest.parent.mkdir(parents=True, exist_ok=True)
        target_gold.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, dest)
        shutil.copyfile(gold, target_gold)
        add(f'POLLOCK-{index:03d}', dest, 'pollock', target_gold)
    # These records are fixed and saved before any format is generated.
    records = [dict(code='00123', count=123 + i, empty=None, korean=f'한글 정산 {i:02d}', note='인용 "확인"\n다음 줄', active=True)
               for i in range(10)]
    write_json(output / 'gold/preservation-records.json', records)
    from openpyxl import Workbook
    import pymupdf as fitz
    formats = ['csv', 'xlsx', 'json', 'txt', 'md', 'pdf']
    for index, record in enumerate(records, 1):
        rows = [record, {**record, 'code': '00007', 'active': False}]
        prose = '\n'.join(f'{key}: {value}' for row in rows for key, value in row.items()) + '\n'
        for fmt in formats:
            source = output / 'inputs/preservation' / f'preserve-{index:02d}.{fmt}'
            source.parent.mkdir(parents=True, exist_ok=True)
            string_rows = [{k: '' if v is None else str(v) for k, v in row.items()} for row in rows]
            expected = dict(rows=string_rows if fmt == 'csv' else rows, columns=list(record), text=prose,
                            tokens=['00123', '00007', str(record['count']), record['korean'], '인용', '다음 줄'])
            if fmt == 'xlsx':
                expected['structure'] = dict(sheet_names=['Data', 'Audit'], formula='=SUM(Data!B2:B3)', merged='A4:B4', audit_string='00099')
            gold = output / 'gold/preservation' / f'preserve-{index:02d}-{fmt}.json'
            write_json(gold, expected)
            if fmt == 'csv':
                with source.open('w', encoding='utf-8-sig', newline='') as stream:
                    writer = csv.DictWriter(stream, fieldnames=list(record)); writer.writeheader(); writer.writerows(string_rows)
            elif fmt == 'json':
                write_json(source, rows)
            elif fmt == 'xlsx':
                book = Workbook(); sheet = book.active; sheet.title = 'Data'
                sheet.append(list(record))
                for row in rows: sheet.append(list(row.values()))
                audit = book.create_sheet('Audit'); audit['A1'] = '원장'; audit['A2'] = '00099'; audit['B2'] = '=SUM(Data!B2:B3)'
                audit.merge_cells('A4:B4'); audit['A4'] = '병합값'; book.save(source)
            elif fmt == 'pdf':
                with fitz.open() as document:
                    page = document.new_page()
                    for row_index, line in enumerate(prose.splitlines()):
                        page.insert_text((24, 30 + row_index * 18), line, fontname='korea', fontsize=11)
                    document.save(source)
            else:
                source.write_text(prose, encoding='utf-8')
            add(f'PRESERVE-{fmt.upper()}-{index:02d}', source, 'preservation', gold)
    corrupt = output / 'inputs/corrupt/broken.pdf'
    corrupt.parent.mkdir(parents=True, exist_ok=True); corrupt.write_bytes(b'%PDF-1.7\nintentionally truncated fixture\n')
    add('CORRUPT-PDF', corrupt, 'robustness', expected_failure=True)
    # Bundle IDs name occurrences; the same corpus file can occur in different runs.
    bundles = []
    for fmt in formats:
        ids = [c['id'] for c in cases if c['suite'] == 'preservation' and c['format'] == fmt]
        for count in [5, 10]:
            bundles.append(dict(id=f'{fmt}-{count}', cases=[dict(corpus_id=identifier) for identifier in ids[:count]]))
    mixed = [f'PRESERVE-{fmt.upper()}-01' for fmt in ['pdf', 'xlsx', 'csv', 'json', 'txt', 'md']]
    bundles.extend([dict(id='mixed', cases=[dict(corpus_id=x) for x in mixed]),
                    dict(id='mixed-reversed', cases=[dict(corpus_id=x) for x in reversed(mixed)]),
                    dict(id='same-name', cases=[dict(corpus_id=f'PRESERVE-TXT-0{i}', upload_name='same-name.txt') for i in [1, 2]]),
                    dict(id='corrupt-between-valid', cases=[dict(corpus_id=x) for x in ['PRESERVE-TXT-01', 'CORRUPT-PDF', 'PRESERVE-TXT-02']])])
    manifest = dict(version='a4b.document_benchmark.v1', seed=SEED, gold_frozen_before_inference=True,
                    dataset_revisions=dict(omni_data=REVISION, omni_code='337cc26965893db3ef53ddc119a6d6bb5bde096f', pollock='bd2a646f4cf55d2bd4e76df511c9c9f675ae7bcc'),
                    omni_population=len(gold_all), omni_pdf_eligible={k: len(v) for k, v in sorted(groups.items())}, cases=cases, bundles=bundles)
    write_json(manifest_path, manifest)
    print(json.dumps(dict(manifest=str(manifest_path), sha256=digest(manifest_path), cases=len(cases), bundles=len(bundles)), ensure_ascii=False))


if __name__ == '__main__':
    main()
